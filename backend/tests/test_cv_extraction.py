import io
import zipfile

import pytest
from docx import Document

from app.config import settings
from app.services.ai_service import ai_service
from app.i18n import t
from app.services.upload_service import (
    NO_TEXT_PDF_MESSAGE,
    UNREADABLE_PDF_MESSAGE,
    ExtractionError,
    PDFExtractionError,
    UploadService,
)
from tests.docx_factory import make_docx
from tests.pdf_factory import make_pdf

MISMATCH_MESSAGE = t("upload.signature_mismatch")
UNREADABLE_DOCX_MESSAGE = t("upload.docx_unreadable")

service = UploadService()


# --- Extraction ---------------------------------------------------------------


def test_extracts_text_from_pdf_bytes():
    assert "Jean Dupont" in service.extract_text_from_bytes(make_pdf("Jean Dupont Python"))


def test_pdf_without_text_returns_empty_string():
    assert service.extract_text_from_bytes(make_pdf(None)) == ""


def test_corrupted_pdf_raises_clear_error():
    with pytest.raises(PDFExtractionError, match="PDF illisible"):
        service.extract_text_from_bytes(b"%PDF-1.4 not really a pdf")


def test_extract_cv_text_handles_pdf_and_txt():
    assert "Jean Dupont" in service.extract_cv_text("CV.PDF", make_pdf("Jean Dupont"))
    assert service.extract_cv_text("cv.txt", "Développeuse Vue".encode("utf-8")) == "Développeuse Vue"
    assert service.extract_cv_text("cv.txt", "Développeuse".encode("latin-1")) == "Développeuse"


def test_extract_cv_text_handles_docx():
    content = make_docx(paragraphs=["Jean Dupont", "Développeur Python"])
    assert service.extract_cv_text("CV.DOCX", content) == "Jean Dupont\nDéveloppeur Python"


def test_docx_extraction_keeps_lists_and_tables_in_order():
    content = make_docx(
        paragraphs=["Expériences"],
        bullets=["Vue 3", "FastAPI"],
        table=[["2022", "Acme"], ["2020", "Globex"]],
    )
    assert service.extract_text_from_docx(content).splitlines() == [
        "Expériences",
        "- Vue 3",
        "- FastAPI",
        "2022 | Acme",
        "2020 | Globex",
    ]


def test_docx_merged_cells_are_not_duplicated():
    document = Document()
    table = document.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Compétences"
    table.cell(0, 2).text = "Python"
    out = io.BytesIO()
    document.save(out)
    assert service.extract_text_from_docx(out.getvalue()) == "Compétences | Python"


@pytest.mark.parametrize(
    "filename,content,code",
    [
        ("cv.pdf", b"garbage", "upload.signature_mismatch"),
        ("cv.pdf", make_docx(paragraphs=["x"]), "upload.signature_mismatch"),
        ("cv.docx", make_pdf("Jean Dupont"), "upload.signature_mismatch"),
        ("cv.docx", b"PK\x03\x04 not a zip", "upload.docx_unreadable"),
        ("cv.txt", make_pdf("Jean Dupont"), "upload.signature_mismatch"),
        ("cv.txt", "Jean".encode("utf-16"), "upload.signature_mismatch"),
    ],
)
def test_content_must_match_extension(filename, content, code):
    with pytest.raises(ExtractionError) as error:
        service.extract_cv_text(filename, content)
    assert error.value.code == code


def test_zip_without_word_document_is_rejected():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ExtractionError) as error:
        service.extract_cv_text("cv.docx", out.getvalue())
    assert error.value.code == "upload.signature_mismatch"


def test_docx_zip_bomb_is_rejected(monkeypatch):
    monkeypatch.setattr("app.services.upload_service.MAX_DOCX_UNCOMPRESSED_SIZE", 1000)
    content = make_docx(paragraphs=["x" * 5000])
    with pytest.raises(ExtractionError) as error:
        service.extract_cv_text("cv.docx", content)
    assert error.value.code == "upload.docx_unreadable"


# --- Upload CV ----------------------------------------------------------------


def _upload(client, headers, content, filename="cv.pdf"):
    return client.post("/api/upload-cv", files={"file": (filename, content, "application/pdf")}, headers=headers)


def test_upload_cv_extracts_pdf(client, auth_headers):
    body = _upload(client, auth_headers, make_pdf("Jean Dupont")).json()
    assert body["success"] is True
    assert "Jean Dupont" in body["text"]


@pytest.mark.parametrize(
    "content,message",
    [
        (make_pdf(None), NO_TEXT_PDF_MESSAGE),
        (b"%PDF-1.4 garbage", UNREADABLE_PDF_MESSAGE),
        (b"garbage", MISMATCH_MESSAGE),
    ],
)
def test_upload_cv_returns_clear_message_for_unreadable_pdf(client, auth_headers, content, message):
    body = _upload(client, auth_headers, content).json()
    assert body == {"success": False, "text": "", "message": message}


def test_upload_cv_extracts_docx_and_txt(client, auth_headers):
    body = _upload(client, auth_headers, make_docx(paragraphs=["Jean Dupont"]), "cv.docx").json()
    assert body["success"] is True
    assert body["text"] == "Jean Dupont"

    body = _upload(client, auth_headers, "Développeuse Vue".encode(), "cv.txt").json()
    assert body["text"] == "Développeuse Vue"


@pytest.mark.parametrize(
    "filename,content,message",
    [
        ("cv.docx", make_docx(), t("upload.empty_file")),
        ("cv.docx", b"PK\x03\x04 broken", UNREADABLE_DOCX_MESSAGE),
        ("cv.txt", b"   ", t("upload.empty_file")),
    ],
)
def test_upload_cv_returns_clear_message_for_unreadable_docx_or_txt(
    client, auth_headers, filename, content, message
):
    body = _upload(client, auth_headers, content, filename).json()
    assert body == {"success": False, "text": "", "message": message}


def test_upload_cv_rejects_unsupported_format(client, auth_headers):
    response = _upload(client, auth_headers, b"{\\rtf1}", "cv.doc")
    assert response.status_code == 400
    assert response.json()["code"] == "upload.cv_format"


def test_upload_cv_rejects_too_large_file(client, auth_headers, monkeypatch):
    monkeypatch.setattr(settings, "MAX_FILE_SIZE", 100)
    response = _upload(client, auth_headers, make_pdf("x" * 200))
    assert response.status_code == 413


def test_free_upload_cv_uses_same_extraction(client):
    response = client.post(
        "/api/free-upload-cv", files={"file": ("cv.pdf", make_pdf("Jean Dupont"), "application/pdf")}
    )
    assert "Jean Dupont" in response.json()["text"]


# --- Interview ----------------------------------------------------------------


@pytest.fixture
def captured_cv(monkeypatch):
    captured = {}

    def fake_generate(cv_text, job_text, num_questions):
        captured["cv_text"] = cv_text
        return [{"id": 1, "question": "Parlez-moi de vous"}]

    monkeypatch.setattr(ai_service, "generate_interview_questions", fake_generate)
    return captured


def _generate(client, headers, content, filename):
    return client.post(
        "/api/interview/generate-questions",
        files={"cv_file": (filename, content)},
        data={"job_text": "Développeur Python", "num_questions": "1"},
        headers=headers,
    )


@pytest.mark.parametrize("path", ["/api/interview/generate-questions", "/api/interview/analyze-responses"])
def test_interview_routes_require_auth(client, path):
    assert client.post(path).status_code in (401, 403)


def test_interview_extracts_pdf_text_not_raw_bytes(client, auth_headers, captured_cv):
    response = _generate(client, auth_headers, make_pdf("Jean Dupont Python"), "cv.pdf")
    assert response.status_code == 200, response.text
    assert "Jean Dupont" in captured_cv["cv_text"]
    assert "%PDF" not in captured_cv["cv_text"]


def test_interview_accepts_docx(client, auth_headers, captured_cv):
    response = _generate(client, auth_headers, make_docx(paragraphs=["Jean Dupont"]), "cv.docx")
    assert response.status_code == 200, response.text
    assert captured_cv["cv_text"] == "Jean Dupont"


def test_interview_accepts_txt(client, auth_headers, captured_cv):
    response = _generate(client, auth_headers, "Jean Dupont".encode(), "cv.txt")
    assert response.status_code == 200
    assert captured_cv["cv_text"] == "Jean Dupont"


@pytest.mark.parametrize(
    "content,detail",
    [
        (make_pdf(None), NO_TEXT_PDF_MESSAGE),
        (b"%PDF-1.4 garbage", UNREADABLE_PDF_MESSAGE),
        (b"garbage", MISMATCH_MESSAGE),
    ],
)
def test_interview_rejects_unreadable_pdf(client, auth_headers, captured_cv, content, detail):
    response = _generate(client, auth_headers, content, "cv.pdf")
    assert response.status_code == 400
    assert response.json()["detail"] == detail
    assert "cv_text" not in captured_cv


def test_interview_rejects_too_large_file(client, auth_headers, captured_cv, monkeypatch):
    monkeypatch.setattr(settings, "MAX_FILE_SIZE", 100)
    response = _generate(client, auth_headers, b"x" * 200, "cv.txt")
    assert response.status_code == 413
    assert "cv_text" not in captured_cv
