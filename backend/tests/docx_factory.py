import io

from docx import Document


def make_docx(
    paragraphs: list[str] = (),
    bullets: list[str] = (),
    table: list[list[str]] | None = None,
) -> bytes:
    """DOCX minimal : paragraphes, puis liste à puces, puis tableau (dans cet ordre)."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for text in bullets:
        document.add_paragraph(text, style="List Bullet")
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table):
            for cell, value in zip(row.cells, values):
                cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
